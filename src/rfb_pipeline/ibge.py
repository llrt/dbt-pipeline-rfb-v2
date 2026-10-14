"""Ingestão das fontes gratuitas do IBGE (incremento `enriquecimento_ibge`).

Tabelas (`esquemas.TABELAS_IBGE`), todas em `raw/ibge/<tabela>/<tabela>.parquet`, all-VARCHAR como
publicadas (ADR-0002), mais `_arquivo_origem` e `_ingerido_em`; fontes anuais ou estáticas, sem
partição por mês:

- `cempre_categorias_cnae`: API SIDRA, tabela 9528 (CEMPRE), nível Brasil, variável 706, todas as
  categorias da classificação 12762 (CNAE 2.0) com códigos **e nomes** (`/f/a`). Dá o período
  publicado mais recente e o nome de cada categoria (`47.41-5 Comércio varejista de tintas...`): é
  por ele que se escolhem as classes (nível 4 da CNAE) e se obtém o código da classe no staging.
- `cempre_municipio_cnae`: tabela 9528 por município (N6), variáveis 706 (unidades locais), 707
  (pessoal ocupado total), 708 (pessoal ocupado assalariado) e 662 (salários, mil reais), para as
  673 classes CNAE e o Total (117897), num único período (`RFB_CEMPRE_PERIODO`, padrão o mais
  recente). Só códigos (`/f/c`).
- `renda_censo_2022_municipio`: tabela 10295 (Censo 2022) por município, variáveis 13604
  (moradores), 13431 (rendimento domiciliar per capita médio) e 13534 (mediano), nos totais das
  classificações 2, 86 e 58.
- `regic_2018_hierarquia`, `regic_2018_ligacoes`, `regic_2018_atracao`: abas das planilhas da REGIC
  2018 no geoftp do IBGE, com todas as colunas; os cabeçalhos viram identificadores (`codmun`,
  `regiao_de_influencia_1`, `q1`...) e o cabeçalho publicado fica documentado na fonte dbt
  (`meta.campo_ibge`).

Colunas SIDRA (`/h/n`, sem a linha de cabeçalho): `nc` (nível territorial), `mc` (unidade de
medida), `v` (valor, **texto** com os símbolos do IBGE: `-` zero absoluto, `X` sigilo, `..` não se
aplica, `...` não disponível), `d1c`... (códigos das dimensões, na ordem da URL: localidade,
variável, período e classificações); com `/f/a`, também `nn`, `mn`, `d1n`... (nomes).

API SIDRA: no máximo 50.000 valores por requisição; acima disso ela responde texto ("excedeu o
limite"), que vira `SidraLimiteExcedidoErro`. O CEMPRE municipal é fatiado por categoria (2 por
requisição, ~340 requisições) e cada fatia vai para `_baixados/ibge/sidra/<tabela>/<período>/` assim
que chega: uma execução interrompida retoma sem refazer as fatias prontas.

Atualização: `rfb ingerir` não baixa de novo o que já está no raw para o mesmo período (SIDRA) ou
arquivo (REGIC), conforme `_manifestos/ibge/manifesto.json`; a renovação é explícita
(`--renovar-ibge`). Só o CEMPRE consulta a SIDRA para saber se saiu um ano novo (uma requisição
leve de categorias), no máximo uma vez a cada 30 dias (`sonda_cempre_em` no manifesto IBGE).

Falha do IBGE não derruba o mês (RIBG-02): a tabela que falhar (rede, API, planilha) fica com o raw
anterior, ou vazia se não havia, a falha vai para `falhas` no manifesto IBGE e a ingestão segue com
aviso. `rfb ingerir --sem-ibge` nem consulta o IBGE. Antes de toda chamada ao dbt, o raw IBGE que
faltar é publicado vazio (`garantir_raw_ibge`, RIBG-01).

`--origem-local DIR` lê, em vez da rede, `DIR/ibge/`:

- `cempre_categorias_cnae.json`, `cempre_municipio_cnae.json`, `renda_censo_2022_municipio.json`:
  respostas da SIDRA no formato acima (lista JSON de objetos; categorias com `/f/a`, as demais com
  `/f/c`), todas as linhas num único arquivo;
- `REGIC2018_Municipios_Hierarquia_e_regiao.xlsx` e
  `REGIC2018_Municipios_Ligacoes_e_atracao_xlsx.zip`, como publicados no geoftp.

Com origem local, uma tabela só é reconvertida se o sha256 do arquivo mudar. Arquivo ausente (por
exemplo, um espelho offline anterior ao incremento, sem `ibge/`) não é erro: a ingestão avisa e
segue, e a tabela fica no raw só com o esquema (Parquet vazio), salvo se já existir lá. O dbt
constrói tudo com as colunas do IBGE NULL (compatibilidade da v1.4 com a v1.3, docs/OPERACAO.md).
"""

from __future__ import annotations

import csv
import hashlib
import json
import math
import re
import shutil
import sys
import time
import unicodedata
import urllib.parse
import uuid
from collections.abc import Callable, Iterator, Sequence
from dataclasses import asdict, dataclass
from datetime import date, datetime, timedelta
from pathlib import Path

import duckdb
import httpx
import openpyxl

from rfb_pipeline.cliente_rfb import _recuo_s, _verificar_host_permitido, baixar_com_retentativas
from rfb_pipeline.configuracao import Configuracao
from rfb_pipeline.conversao import (
    _conectar,
    _dir_temporario,
    _literal,
    _publicar,
    _ts,
    extrair_zip_seguro,
)
from rfb_pipeline.erros import (
    ConversaoErro,
    EntidadeVaziaErro,
    ErroIngestao,
    PlanilhaIBGEErro,
    SidraErro,
    SidraLimiteExcedidoErro,
)
from rfb_pipeline.esquemas import TABELAS_IBGE
from rfb_pipeline.manifesto import sha256_arquivo

__all__ = [
    "ConsultaSidra",
    "EntradaIBGE",
    "consultar_sidra",
    "fatiar_categorias",
    "garantir_raw_ibge",
    "ingerir_ibge",
    "ler_aba_xlsx",
    "ler_manifesto_ibge",
    "normalizar_coluna",
]

# ------------------------------------------------------------------------- constantes

CATEGORIA_TOTAL_CEMPRE = "117897"
RE_CLASSE_CNAE = re.compile(r"^\d{2}\.\d{2}-\d ")  # nome de classe: `47.41-5 Comércio...`
MUNICIPIOS_ESTIMADOS = 5_600  # teto do fatiamento (o Brasil tem 5.570 municípios no N6)
# RIBG-02: a consulta de ano novo do CEMPRE (categorias, nível Brasil) roda no máximo uma vez neste
# intervalo; `--renovar-ibge` consulta sempre
INTERVALO_SONDA_CEMPRE = timedelta(days=30)

_TABELA_CEMPRE = "9528"
_VARIAVEIS_CEMPRE = ("706", "707", "708", "662")
_CLASSIFICACAO_CEMPRE = "12762"
_TABELA_RENDA = "10295"
_VARIAVEIS_RENDA = ("13604", "13431", "13534")
_CLASSIFICACOES_RENDA = (("2", ("6794",)), ("86", ("95251",)), ("58", ("95253",)))
_PERIODO_RENDA = "2022"

_ARQUIVO_HIERARQUIA = "REGIC2018_Municipios_Hierarquia_e_regiao.xlsx"
_ARQUIVO_LIGACOES_ATRACAO = "REGIC2018_Municipios_Ligacoes_e_atracao_xlsx.zip"


@dataclass(frozen=True)
class PlanilhaRegic:
    tabela: str
    arquivo: str  # nome do arquivo no geoftp (xlsx ou zip)
    membro_zip: str | None  # xlsx dentro do zip, se houver
    aba: str


PLANILHAS_REGIC: tuple[PlanilhaRegic, ...] = (
    PlanilhaRegic("regic_2018_hierarquia", _ARQUIVO_HIERARQUIA, None, "Hierarquia e região"),
    PlanilhaRegic(
        "regic_2018_ligacoes",
        _ARQUIVO_LIGACOES_ATRACAO,
        "REGIC2018_Quest_Ligacoes_entre_Municipios.xlsx",
        "BD_LIG_MUN_QUEST",
    ),
    PlanilhaRegic(
        "regic_2018_atracao",
        _ARQUIVO_LIGACOES_ATRACAO,
        "REGIC2018_Quest_Atracao_Municipios.xlsx",
        "BD_MUN_QUEST",
    ),
)
_PERIODO_REGIC = "2018"

# --------------------------------------------------------------------------- SIDRA


@dataclass(frozen=True)
class ConsultaSidra:
    """Uma requisição à API SIDRA (`/values`), sempre sem cabeçalho (`/h/n`)."""

    tabela: str
    nivel: str  # "n6" (município) | "n1" (Brasil)
    variaveis: tuple[str, ...]
    periodo: str  # ano ou "last 1"
    classificacoes: tuple[tuple[str, tuple[str, ...]], ...]  # (id, categorias) ou ("x", ("all",))
    formato: str = "c"  # "c": só códigos; "a": códigos e nomes
    localidades: str = "all"

    def url(self, base_url: str) -> str:
        partes = [
            base_url.rstrip("/"),
            "t",
            self.tabela,
            self.nivel,
            self.localidades,
            "v",
            ",".join(self.variaveis),
            "p",
            urllib.parse.quote(self.periodo),
        ]
        for classificacao, categorias in self.classificacoes:
            partes += [f"c{classificacao}", ",".join(categorias)]
        partes += ["f", self.formato, "h", "n"]
        return "/".join(partes)

    def valores_estimados(self, localidades: int) -> int:
        """Teto de valores pedidos (localidades × variáveis × categorias de cada classificação)."""
        total = localidades * len(self.variaveis)
        for _, categorias in self.classificacoes:
            total *= len(categorias)
        return total


def fatiar_categorias(
    categorias: Sequence[str], valores_por_categoria: int, limite: int
) -> list[tuple[str, ...]]:
    """Agrupa `categorias` em fatias de até `limite` valores (`len × valores_por_categoria`)."""
    por_fatia = limite // valores_por_categoria
    if por_fatia < 1:
        raise ValueError(
            f"uma categoria já pede {valores_por_categoria} valores, acima do limite {limite}"
        )
    return [tuple(categorias[i : i + por_fatia]) for i in range(0, len(categorias), por_fatia)]


def _validar_resposta_sidra(url: str, conteudo: bytes) -> list[dict]:
    """Corpo de uma resposta 200: lista JSON não vazia de objetos com as mesmas chaves."""
    try:
        dados = json.loads(conteudo)
    except ValueError:
        trecho = conteudo[:300].decode("utf-8", errors="replace").strip()
        if "excedeu o limite" in trecho:
            raise SidraLimiteExcedidoErro(url, trecho) from None
        raise SidraErro(url, f"resposta não é JSON: {trecho!r}") from None
    if not isinstance(dados, list) or not dados:
        raise SidraErro(url, "resposta vazia (período, variável ou categoria inexistente?)")
    chaves = list(dados[0]) if isinstance(dados[0], dict) else None
    if not chaves or any(not isinstance(d, dict) or list(d) != chaves for d in dados):
        raise SidraErro(url, "resposta com objetos de chaves diferentes")
    if chaves[0] == "NC" and dados[0].get("NC") == "Nível Territorial (Código)":
        raise SidraErro(url, "resposta com linha de cabeçalho (use /h/n)")
    return dados


def consultar_sidra(
    http: httpx.Client,
    url: str,
    *,
    hosts_permitidos: tuple[str, ...],
    tentativas: int,
    dormir: Callable[[float], None] = time.sleep,
) -> tuple[bytes, list[dict]]:
    """GET na SIDRA com retentativa e recuo (rede, 5xx, 408/429), como `baixar_com_retentativas`.

    Erros da API vêm como texto (HTTP 400): "excedeu o limite" levanta `SidraLimiteExcedidoErro`;
    os demais, `SidraErro` com o texto. Devolve o corpo e a lista de objetos já validada.
    """
    _verificar_host_permitido(url, hosts_permitidos)
    tentativas = max(tentativas, 1)
    for numero in range(1, tentativas + 1):
        try:
            resposta = http.get(url)
        except httpx.TransportError as exc:
            if numero == tentativas:
                raise SidraErro(url, f"falha de rede após {numero} tentativa(s): {exc}") from exc
            dormir(_recuo_s(numero))
            continue
        codigo = resposta.status_code
        if codigo >= 500 or codigo in (408, 429):
            if numero == tentativas:
                raise SidraErro(url, f"HTTP {codigo} após {numero} tentativa(s)")
            dormir(_recuo_s(numero))
            continue
        if codigo != 200:
            texto = resposta.text.strip()[:300]
            if "excedeu o limite" in texto:
                raise SidraLimiteExcedidoErro(url, texto)
            raise SidraErro(url, f"HTTP {codigo}: {texto}")
        return resposta.content, _validar_resposta_sidra(url, resposta.content)
    raise AssertionError("inalcançável")  # pragma: no cover


def _sql_json_para_parquet(
    json_path: Path, chaves: Sequence[str], destino: Path, origem: str, ingerido_em: datetime
) -> str:
    """`COPY` de uma resposta SIDRA (lista JSON) para Parquet, todas as colunas VARCHAR."""
    colunas = "{" + ", ".join(f"{_literal(c)}: 'VARCHAR'" for c in chaves) + "}"
    selecao = ", ".join(f'"{c}" AS {c.lower()}' for c in chaves)
    return (
        f"COPY (SELECT {selecao}, {_literal(origem)}::VARCHAR AS _arquivo_origem, "
        f"{_literal(_ts(ingerido_em))}::TIMESTAMP AS _ingerido_em "
        f"FROM read_json({_literal(str(json_path))}, format='array', columns={colunas})) "
        f"TO {_literal(str(destino))} (FORMAT parquet, COMPRESSION zstd)"
    )


def _gravar_json_parquet(
    con: duckdb.DuckDBPyConnection,
    json_path: Path,
    chaves: Sequence[str],
    destino: Path,
    origem: str,
    ingerido_em: datetime,
) -> int:
    try:
        linha = con.execute(
            _sql_json_para_parquet(json_path, chaves, destino, origem, ingerido_em)
        ).fetchone()
    except duckdb.Error as exc:
        raise ConversaoErro(json_path.name, str(exc)) from exc
    return int(linha[0]) if linha else 0


# ------------------------------------------------------------------------ planilhas


def normalizar_coluna(nome: str) -> str:
    """Cabeçalho publicado -> identificador: sem acento, minúsculo, `_` no lugar do resto."""
    ascii_ = unicodedata.normalize("NFKD", nome).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^a-z0-9]+", "_", ascii_.lower()).strip("_")


def _texto_celula(valor: object) -> str | None:
    """Valor de célula como texto: inteiro sem `.0`, decimal na menor representação exata."""
    if valor is None:
        return None
    if isinstance(valor, bool):
        return "true" if valor else "false"
    if isinstance(valor, int):
        return str(valor)
    if isinstance(valor, float):
        if math.isfinite(valor) and valor.is_integer():
            return str(int(valor))
        return repr(valor)
    if isinstance(valor, datetime | date):
        return valor.isoformat()
    return str(valor)


def ler_aba_xlsx(
    caminho: Path, aba: str
) -> tuple[list[str], list[str], Iterator[list[str | None]]]:
    """Lê `aba` de um xlsx (openpyxl, somente leitura): colunas normalizadas, cabeçalho e linhas.

    A primeira linha é o cabeçalho; colunas vazias à direita são ignoradas e linhas totalmente
    vazias, descartadas. Cabeçalho ausente, vazio no meio ou repetido após normalizar levanta
    `PlanilhaIBGEErro`, assim como valor numa coluna sem cabeçalho.
    """
    livro = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    if aba not in livro.sheetnames:
        livro.close()
        raise PlanilhaIBGEErro(caminho.name, f"aba {aba!r} ausente (abas: {livro.sheetnames})")
    linhas = livro[aba].iter_rows(values_only=True)
    cabecalho_bruto = list(next(linhas, ()))
    while cabecalho_bruto and cabecalho_bruto[-1] in (None, ""):
        cabecalho_bruto.pop()
    if not cabecalho_bruto or any(c in (None, "") for c in cabecalho_bruto):
        livro.close()
        raise PlanilhaIBGEErro(caminho.name, f"aba {aba!r} sem cabeçalho completo")
    cabecalho = [str(c).strip() for c in cabecalho_bruto]
    colunas = [normalizar_coluna(c) for c in cabecalho]
    repetidas = sorted({c for c in colunas if colunas.count(c) > 1 or not c})
    if repetidas:
        livro.close()
        raise PlanilhaIBGEErro(caminho.name, f"colunas repetidas ou vazias: {repetidas}")
    largura = len(colunas)

    def _linhas() -> Iterator[list[str | None]]:
        try:
            for numero, linha in enumerate(linhas, start=2):
                valores = [_texto_celula(v) for v in linha]
                if any(v not in (None, "") for v in valores[largura:]):
                    raise PlanilhaIBGEErro(
                        caminho.name, f"aba {aba!r}, linha {numero}: valor sem cabeçalho"
                    )
                valores = (valores + [None] * largura)[:largura]
                if all(v in (None, "") for v in valores):
                    continue
                yield valores
        finally:
            livro.close()

    return colunas, cabecalho, _linhas()


def _gravar_aba_parquet(
    con: duckdb.DuckDBPyConnection,
    xlsx: Path,
    aba: str,
    destino: Path,
    trabalho: Path,
    ingerido_em: datetime,
) -> int:
    """Aba do xlsx -> CSV temporário -> Parquet all-VARCHAR (vazio = NULL, como no resto do raw)."""
    colunas, _, linhas = ler_aba_xlsx(xlsx, aba)
    csv_tmp = trabalho / f"{destino.stem}.csv"
    with csv_tmp.open("w", encoding="utf-8", newline="") as fh:
        escritor = csv.writer(fh, quoting=csv.QUOTE_MINIMAL)
        escritor.writerow(colunas)
        for linha in linhas:
            escritor.writerow(["" if v is None else v for v in linha])
    tipos = "{" + ", ".join(f"{_literal(c)}: 'VARCHAR'" for c in colunas) + "}"
    sql = (
        f"COPY (SELECT *, {_literal(xlsx.name)}::VARCHAR AS _arquivo_origem, "
        f"{_literal(_ts(ingerido_em))}::TIMESTAMP AS _ingerido_em "
        f"FROM read_csv({_literal(str(csv_tmp))}, header=true, delim=',', quote='\"', "
        f"escape='\"', encoding='utf-8', columns={tipos})) "
        f"TO {_literal(str(destino))} (FORMAT parquet, COMPRESSION zstd)"
    )
    try:
        linha = con.execute(sql).fetchone()
    except duckdb.Error as exc:
        raise ConversaoErro(f"{xlsx.name}[{aba}]", str(exc)) from exc
    finally:
        csv_tmp.unlink(missing_ok=True)
    return int(linha[0]) if linha else 0


# ------------------------------------------------------------------------ manifesto


@dataclass(frozen=True)
class EntradaIBGE:
    """Uma tabela IBGE no manifesto (`_manifestos/ibge/manifesto.json` e manifesto do mês)."""

    tabela: str
    origem: str  # URL(s) de origem ou caminho do arquivo local
    periodo: str  # ano publicado (CEMPRE: o baixado; renda: 2022; REGIC: 2018)
    arquivo: str | None  # arquivo baixado/copiado (REGIC e origem local); None na SIDRA remota
    bytes: int  # tamanho do arquivo, ou soma dos corpos das respostas SIDRA
    sha256: str | None  # do arquivo, quando há arquivo
    requisicoes: int  # requisições SIDRA (0 nas planilhas e na origem local)
    linhas: int
    parquet: str  # relativo a RAIZ_DADOS
    ingerido_em: str

    def to_dict(self) -> dict:
        return asdict(self)

    @staticmethod
    def from_dict(dados: dict) -> EntradaIBGE:
        return EntradaIBGE(**dados)


def _caminho_manifesto_ibge(configuracao: Configuracao) -> Path:
    """Fora da raiz de `_manifestos/`, que só guarda um `<AAAA-MM>.json` por mês."""
    return configuracao.manifestos_dir / "ibge" / "manifesto.json"


def _ler_corpo_manifesto_ibge(configuracao: Configuracao) -> dict:
    caminho = _caminho_manifesto_ibge(configuracao)
    if not caminho.exists():
        return {}
    return json.loads(caminho.read_text(encoding="utf-8"))


def ler_manifesto_ibge(configuracao: Configuracao) -> dict[str, EntradaIBGE]:
    dados = _ler_corpo_manifesto_ibge(configuracao)
    return {nome: EntradaIBGE.from_dict(e) for nome, e in dados.get("tabelas", {}).items()}


def ler_falhas_ibge(configuracao: Configuracao) -> dict[str, dict]:
    """Tabelas cuja última tentativa falhou (`{tabela: {"erro", "em"}}`), do manifesto IBGE."""
    return dict(_ler_corpo_manifesto_ibge(configuracao).get("falhas", {}))


def _escrever_manifesto_ibge(
    configuracao: Configuracao,
    entradas: dict[str, EntradaIBGE],
    *,
    sonda_cempre_em: str | None = None,
    falhas: dict[str, dict] | None = None,
) -> None:
    """Chaves opcionais (RIBG-02): `sonda_cempre_em`, a última consulta de ano novo do CEMPRE, e
    `falhas`, as tabelas que não foram atualizadas na última tentativa (raw anterior mantido)."""
    destino = _caminho_manifesto_ibge(configuracao)
    destino.parent.mkdir(parents=True, exist_ok=True)
    tmp = destino.with_name(f".tmp-{destino.name}-{uuid.uuid4().hex}")
    corpo: dict = {"tabelas": {nome: e.to_dict() for nome, e in sorted(entradas.items())}}
    if sonda_cempre_em is not None:
        corpo["sonda_cempre_em"] = sonda_cempre_em
    if falhas:
        corpo["falhas"] = dict(sorted(falhas.items()))
    tmp.write_text(json.dumps(corpo, ensure_ascii=False, indent=2), encoding="utf-8")
    tmp.replace(destino)


def _parquet_final(configuracao: Configuracao, tabela: str) -> Path:
    return configuracao.raw_dir / "ibge" / tabela / f"{tabela}.parquet"


def _publicar_vazia(configuracao: Configuracao, tabela: str) -> None:
    """Grava `raw/ibge/<tabela>/<tabela>.parquet` só com o cabeçalho (all-VARCHAR, ADR-0002)."""
    final = _parquet_final(configuracao, tabela)
    final.parent.parent.mkdir(parents=True, exist_ok=True)
    tmp_dir = _dir_temporario(final.parent)
    colunas = ", ".join(f"NULL::VARCHAR AS {coluna}" for coluna in TABELAS_IBGE[tabela].colunas)
    try:
        with duckdb.connect() as con:
            con.execute(
                f"COPY (SELECT {colunas}, NULL::VARCHAR AS _arquivo_origem, "
                f"NULL::TIMESTAMP AS _ingerido_em LIMIT 0) "
                f"TO {_literal(str(tmp_dir / final.name))} (FORMAT parquet)"
            )
        _publicar(tmp_dir, final.parent)
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)


def garantir_raw_ibge(configuracao: Configuracao) -> list[str]:
    """Publica vazia (só o esquema) cada tabela IBGE que falta no raw; devolve as publicadas.

    Sem rede e sem tocar num raw existente. As fontes dbt do IBGE leem
    `read_parquet(.../raw/ibge/<tabela>/*.parquet)`, que falha sem arquivo: um `RAIZ_DADOS` da v1.3
    (sem `raw/ibge/`) precisa disto antes de qualquer chamada ao dbt (RIBG-01). Com o raw vazio, o
    dbt constrói tudo com as colunas do IBGE NULL.
    """
    publicadas: list[str] = []
    for tabela in TABELAS_IBGE:
        if not _parquet_final(configuracao, tabela).exists():
            _publicar_vazia(configuracao, tabela)
            publicadas.append(tabela)
    if publicadas:
        print(
            f"ibge: AVISO raw ausente de {', '.join(publicadas)}; publicado vazio (só o esquema), "
            "as colunas do IBGE ficam NULL até uma ingestão (docs/OPERACAO.md)",
            flush=True,
        )
    return publicadas


def _reaproveitavel(
    configuracao: Configuracao,
    anterior: EntradaIBGE | None,
    *,
    periodo: str | None = None,
    sha256: str | None = None,
) -> bool:
    """A tabela já está no raw com o mesmo período (SIDRA) ou o mesmo arquivo (sha256)?"""
    if anterior is None or not _parquet_final(configuracao, anterior.tabela).exists():
        return False
    if periodo is not None and anterior.periodo != periodo:
        return False
    return sha256 is None or anterior.sha256 == sha256


# ------------------------------------------------------------------------ ingestão


class _Ingestao:
    """Estado de uma execução: configuração, HTTP, conexão DuckDB e manifesto em construção."""

    def __init__(
        self,
        configuracao: Configuracao,
        *,
        origem_local: Path | None,
        renovar: bool,
        ingerido_em: datetime,
        http: httpx.Client,
        dormir: Callable[[float], None],
    ) -> None:
        self.configuracao = configuracao
        self.origem_local = origem_local
        self.renovar = renovar
        self.ingerido_em = ingerido_em
        self.http = http
        self.dormir = dormir
        self.anteriores = ler_manifesto_ibge(configuracao)
        self.entradas: dict[str, EntradaIBGE] = dict(self.anteriores)
        corpo = _ler_corpo_manifesto_ibge(configuracao)
        self.sonda_cempre_em: str | None = corpo.get("sonda_cempre_em")
        self.falhas: dict[str, dict] = dict(corpo.get("falhas", {}))
        self.baixados = configuracao.raiz_dados / "_baixados" / "ibge"
        self._regic_obtidos: dict[str, Path] = {}

    # -- utilitários

    def _local(self, nome: str) -> Path:
        assert self.origem_local is not None
        return self.origem_local / "ibge" / nome

    def _ausente_local(self, tabela: str, arquivo: Path) -> bool:
        """Modo offline sem o arquivo da tabela: avisa, garante o raw vazio e devolve True.

        Compatibilidade com espelhos offline anteriores ao incremento (sem `ibge/`): a ingestão
        segue e o dbt constrói tudo, com as colunas do IBGE NULL. Um raw já existente fica como
        está.
        """
        if self.origem_local is None or arquivo.is_file():
            return False
        final = _parquet_final(self.configuracao, tabela)
        if final.exists():
            print(f"ibge.{tabela}: AVISO {arquivo} ausente; mantido o raw existente", flush=True)
            return True
        self._publicar_vazia(tabela)
        print(
            f"ibge.{tabela}: AVISO {arquivo} ausente; publicada vazia (só o esquema), "
            "as colunas do IBGE ficam NULL",
            flush=True,
        )
        return True

    def _publicar_vazia(self, tabela: str) -> None:
        _publicar_vazia(self.configuracao, tabela)

    def _pular(self, tabela: str, motivo: str) -> None:
        entrada = self.anteriores[tabela]
        self.entradas[tabela] = entrada
        if self.falhas.pop(tabela, None) is not None:
            self._gravar_manifesto()
        print(f"ibge.{tabela}: pulada ({motivo}) - {entrada.linhas} linhas", flush=True)

    def _gravar_manifesto(self) -> None:
        _escrever_manifesto_ibge(
            self.configuracao,
            self.entradas,
            sonda_cempre_em=self.sonda_cempre_em,
            falhas=self.falhas,
        )

    def falhou(self, tabela: str, erro: BaseException) -> None:
        """RIBG-02: a tabela não foi atualizada. Avisa, mantém o raw anterior (ou publica o vazio)
        e registra a falha no manifesto IBGE; a ingestão do mês segue."""
        final = _parquet_final(self.configuracao, tabela)
        destino = "mantido o raw anterior" if final.exists() else "publicada vazia (só o esquema)"
        if not final.exists():
            _publicar_vazia(self.configuracao, tabela)
        motivo = str(erro).strip() or type(erro).__name__
        print(
            f"aviso: ibge.{tabela}: não atualizada ({motivo[:300]}); {destino}. A RFB e a BD "
            "seguem; tente de novo com `rfb ingerir --renovar-ibge` (docs/OPERACAO.md)",
            file=sys.stderr,
            flush=True,
        )
        self.falhas[tabela] = {"erro": motivo[:500], "em": self.ingerido_em.isoformat()}
        self._gravar_manifesto()

    def _registrar(self, entrada: EntradaIBGE, inicio: float) -> None:
        self.entradas[entrada.tabela] = entrada
        self.falhas.pop(entrada.tabela, None)
        self._gravar_manifesto()
        print(
            f"ibge.{entrada.tabela}: {entrada.linhas} linhas, {entrada.bytes} bytes, "
            f"{time.monotonic() - inicio:.1f}s",
            flush=True,
        )

    def _publicar_parquet(
        self, tabela: str, gravar: Callable[[duckdb.DuckDBPyConnection, Path, Path], int]
    ) -> tuple[int, str]:
        """Grava `raw/ibge/<tabela>/<tabela>.parquet` atomicamente; 0 linhas não publica.

        `gravar(con, destino, trabalho)` escreve o Parquet em `destino`; arquivos intermediários
        (JSON, CSV, xlsx extraído) vão para `trabalho`, fora do diretório publicado.
        """
        final = _parquet_final(self.configuracao, tabela)
        final.parent.parent.mkdir(parents=True, exist_ok=True)
        tmp_dir = _dir_temporario(final.parent)
        trabalho = self.configuracao.raiz_dados / "_tmp" / f"extracao-ibge-{uuid.uuid4().hex}"
        trabalho.mkdir(parents=True)
        con: duckdb.DuckDBPyConnection | None = None
        try:
            con = _conectar(self.configuracao)
            linhas = gravar(con, tmp_dir / final.name, trabalho)
            if linhas == 0:
                raise EntidadeVaziaErro(tabela, [final.name])
            con.close()
            con = None
            _publicar(tmp_dir, final.parent)
        finally:
            if con is not None:
                con.close()
            shutil.rmtree(tmp_dir, ignore_errors=True)
            shutil.rmtree(trabalho, ignore_errors=True)
        return linhas, str(final.relative_to(self.configuracao.raiz_dados))

    def _sonda_recente(self) -> bool:
        """A consulta de ano novo do CEMPRE foi feita há menos de `INTERVALO_SONDA_CEMPRE`?"""
        if self.sonda_cempre_em is None:
            return False
        try:
            ultima = datetime.fromisoformat(self.sonda_cempre_em)
        except ValueError:
            return False
        return timedelta(0) <= self.ingerido_em - ultima < INTERVALO_SONDA_CEMPRE

    def _consultar(self, url: str) -> tuple[bytes, list[dict]]:
        return consultar_sidra(
            self.http,
            url,
            hosts_permitidos=self.configuracao.hosts_permitidos,
            tentativas=self.configuracao.tentativas,
            dormir=self.dormir,
        )

    # -- SIDRA

    def _sidra_local(self, tabela: str) -> tuple[Path, list[dict], str]:
        arquivo = self._local(f"{tabela}.json")
        dados = _validar_resposta_sidra(str(arquivo), arquivo.read_bytes())
        return arquivo, dados, sha256_arquivo(arquivo)

    def _gravar_sidra_unica(
        self,
        tabela: str,
        conteudo_ou_arquivo: bytes | Path,
        chaves: Sequence[str],
        origem: str,
    ) -> tuple[int, str]:
        def gravar(con: duckdb.DuckDBPyConnection, destino: Path, trabalho: Path) -> int:
            if isinstance(conteudo_ou_arquivo, Path):
                json_path = conteudo_ou_arquivo
            else:
                json_path = trabalho / f"{tabela}.json"
                json_path.write_bytes(conteudo_ou_arquivo)
            return _gravar_json_parquet(con, json_path, chaves, destino, origem, self.ingerido_em)

        return self._publicar_parquet(tabela, gravar)

    def categorias_cempre(self) -> tuple[str | None, list[dict]]:
        """Período do CEMPRE e as categorias CNAE (com nomes); grava `cempre_categorias_cnae`.

        Offline sem o arquivo: `(None, [])` (o CEMPRE municipal offline não usa o período).
        """
        tabela = "cempre_categorias_cnae"
        inicio = time.monotonic()
        if self.origem_local is not None:
            if self._ausente_local(tabela, self._local(f"{tabela}.json")):
                return None, []
            arquivo, dados, sha = self._sidra_local(tabela)
            periodo = _periodo_unico(dados, "D3C", str(arquivo))
            if not self.renovar and _reaproveitavel(
                self.configuracao, self.anteriores.get(tabela), sha256=sha
            ):
                self._pular(tabela, "arquivo local sem mudanças")
                return periodo, dados
            linhas, parquet = self._gravar_sidra_unica(
                tabela, arquivo, list(dados[0]), arquivo.name
            )
            entrada = EntradaIBGE(
                tabela, str(arquivo), periodo, arquivo.name, arquivo.stat().st_size, sha, 0,
                linhas, parquet, self.ingerido_em.isoformat(),
            )  # fmt: skip
            self._registrar(entrada, inicio)
            return periodo, dados

        anterior = self.anteriores.get(tabela)
        if not self.renovar and self._sonda_recente() and anterior is not None:
            periodo = anterior.periodo
            if _reaproveitavel(
                self.configuracao, anterior, periodo=self.configuracao.cempre_periodo or periodo
            ) and _reaproveitavel(
                self.configuracao, self.anteriores.get("cempre_municipio_cnae"), periodo=periodo
            ):
                self._pular(
                    tabela,
                    f"ano novo consultado em {self.sonda_cempre_em[:10]}; próxima consulta "
                    f"{INTERVALO_SONDA_CEMPRE.days} dias depois ou com --renovar-ibge",
                )
                return periodo, []

        consulta = ConsultaSidra(
            tabela=_TABELA_CEMPRE,
            nivel="n1",
            variaveis=("706",),
            periodo=self.configuracao.cempre_periodo or "last 1",
            classificacoes=((_CLASSIFICACAO_CEMPRE, ("all",)),),
            formato="a",
        )
        url = consulta.url(self.configuracao.sidra_base_url)
        conteudo, dados = self._consultar(url)
        periodo = _periodo_unico(dados, "D3C", url)
        self.sonda_cempre_em = self.ingerido_em.isoformat()
        self._gravar_manifesto()
        if not self.renovar and _reaproveitavel(
            self.configuracao, self.anteriores.get(tabela), periodo=periodo
        ):
            self._pular(tabela, f"período {periodo} já no raw")
            return periodo, dados
        linhas, parquet = self._gravar_sidra_unica(tabela, conteudo, list(dados[0]), url)
        entrada = EntradaIBGE(
            tabela, url, periodo, None, len(conteudo), None, 1, linhas, parquet,
            self.ingerido_em.isoformat(),
        )  # fmt: skip
        self._registrar(entrada, inicio)
        return periodo, dados

    def cempre_municipio(self, periodo: str | None, categorias: list[dict]) -> None:
        tabela = "cempre_municipio_cnae"
        inicio = time.monotonic()
        if self.origem_local is not None:
            if self._ausente_local(tabela, self._local(f"{tabela}.json")):
                return
            arquivo, dados, sha = self._sidra_local(tabela)
            if not self.renovar and _reaproveitavel(
                self.configuracao, self.anteriores.get(tabela), sha256=sha
            ):
                self._pular(tabela, "arquivo local sem mudanças")
                return
            periodo_local = _periodo_unico(dados, "D3C", str(arquivo))
            linhas, parquet = self._gravar_sidra_unica(
                tabela, arquivo, list(dados[0]), arquivo.name
            )
            entrada = EntradaIBGE(
                tabela, str(arquivo), periodo_local, arquivo.name, arquivo.stat().st_size, sha,
                0, linhas, parquet, self.ingerido_em.isoformat(),
            )  # fmt: skip
            self._registrar(entrada, inicio)
            return

        assert periodo is not None  # na rede, `categorias_cempre` sempre devolve o período
        if not self.renovar and _reaproveitavel(
            self.configuracao, self.anteriores.get(tabela), periodo=periodo
        ):
            self._pular(tabela, f"período {periodo} já no raw")
            return

        classes = sorted({d["D4C"] for d in categorias if RE_CLASSE_CNAE.match(d.get("D4N", ""))})
        if not classes:
            raise SidraErro("(categorias do CEMPRE)", "nenhuma classe CNAE (`NN.NN-N nome`)")
        escolhidas = [*classes, CATEGORIA_TOTAL_CEMPRE]
        modelo = ConsultaSidra(
            tabela=_TABELA_CEMPRE,
            nivel="n6",
            variaveis=_VARIAVEIS_CEMPRE,
            periodo=periodo,
            classificacoes=((_CLASSIFICACAO_CEMPRE, ("0",)),),
        )
        fatias = fatiar_categorias(
            escolhidas,
            modelo.valores_estimados(MUNICIPIOS_ESTIMADOS),
            self.configuracao.sidra_max_valores,
        )
        pasta = self.baixados / "sidra" / tabela / periodo
        if self.renovar:
            shutil.rmtree(pasta, ignore_errors=True)
        pasta.mkdir(parents=True, exist_ok=True)
        for residuo in pasta.glob(".tmp-*"):
            residuo.unlink(missing_ok=True)

        total_bytes = 0
        reaproveitadas = 0
        esperadas: list[Path] = []  # fatias deste plano (as de outro plano ficam de fora)
        con = _conectar(self.configuracao)
        try:
            for indice, fatia in enumerate(fatias, start=1):
                consulta = ConsultaSidra(
                    tabela=modelo.tabela,
                    nivel=modelo.nivel,
                    variaveis=modelo.variaveis,
                    periodo=periodo,
                    classificacoes=((_CLASSIFICACAO_CEMPRE, fatia),),
                )
                url = consulta.url(self.configuracao.sidra_base_url)
                hash_url = hashlib.sha1(url.encode("utf-8")).hexdigest()[:12]
                destino = pasta / f"fatia-{indice:04d}-{hash_url}.parquet"
                esperadas.append(destino)
                if destino.exists():
                    reaproveitadas += 1
                    total_bytes += int(destino.with_suffix(".bytes").read_text() or 0)
                    continue
                conteudo, dados = self._consultar(url)
                _periodo_unico(dados, "D3C", url, esperado=periodo)
                json_tmp = pasta / f".tmp-{destino.stem}.json"
                parquet_tmp = pasta / f".tmp-{destino.name}"
                json_tmp.write_bytes(conteudo)
                try:
                    _gravar_json_parquet(
                        con, json_tmp, list(dados[0]), parquet_tmp, url, self.ingerido_em
                    )
                finally:
                    json_tmp.unlink(missing_ok=True)
                destino.with_suffix(".bytes").write_text(str(len(conteudo)))
                parquet_tmp.replace(destino)  # a fatia só "existe" depois de completa
                total_bytes += len(conteudo)
                if indice % 10 == 0 or indice == len(fatias):
                    print(
                        f"ibge.{tabela}: fatia {indice}/{len(fatias)} "
                        f"({reaproveitadas} reaproveitadas), {time.monotonic() - inicio:.0f}s",
                        flush=True,
                    )
        finally:
            con.close()

        def gravar(con: duckdb.DuckDBPyConnection, destino: Path, _tmp: Path) -> int:
            arquivos = [str(p) for p in esperadas if p.exists()]
            if len(arquivos) != len(fatias):
                raise SidraErro(
                    "(fatias do CEMPRE)", f"{len(arquivos)} de {len(fatias)} fatias no disco"
                )
            lista = "[" + ", ".join(_literal(a) for a in arquivos) + "]"
            try:
                linha = con.execute(
                    f"COPY (SELECT * FROM read_parquet({lista})) TO {_literal(str(destino))} "
                    "(FORMAT parquet, COMPRESSION zstd)"
                ).fetchone()
            except duckdb.Error as exc:
                raise ConversaoErro(tabela, str(exc)) from exc
            return int(linha[0]) if linha else 0

        linhas, parquet = self._publicar_parquet(tabela, gravar)
        origem = self.configuracao.sidra_base_url + f"/t/{_TABELA_CEMPRE}/n6/all (fatiada)"
        entrada = EntradaIBGE(
            tabela, origem, periodo, None, total_bytes, None, len(fatias), linhas, parquet,
            self.ingerido_em.isoformat(),
        )  # fmt: skip
        self._registrar(entrada, inicio)
        shutil.rmtree(pasta, ignore_errors=True)

    def renda(self) -> None:
        tabela = "renda_censo_2022_municipio"
        inicio = time.monotonic()
        if self.origem_local is not None:
            if self._ausente_local(tabela, self._local(f"{tabela}.json")):
                return
            arquivo, dados, sha = self._sidra_local(tabela)
            if not self.renovar and _reaproveitavel(
                self.configuracao, self.anteriores.get(tabela), sha256=sha
            ):
                self._pular(tabela, "arquivo local sem mudanças")
                return
            periodo = _periodo_unico(dados, "D3C", str(arquivo))
            linhas, parquet = self._gravar_sidra_unica(
                tabela, arquivo, list(dados[0]), arquivo.name
            )
            entrada = EntradaIBGE(
                tabela, str(arquivo), periodo, arquivo.name, arquivo.stat().st_size, sha, 0,
                linhas, parquet, self.ingerido_em.isoformat(),
            )  # fmt: skip
            self._registrar(entrada, inicio)
            return

        if not self.renovar and _reaproveitavel(
            self.configuracao, self.anteriores.get(tabela), periodo=_PERIODO_RENDA
        ):
            self._pular(tabela, f"período {_PERIODO_RENDA} já no raw")
            return
        consulta = ConsultaSidra(
            tabela=_TABELA_RENDA,
            nivel="n6",
            variaveis=_VARIAVEIS_RENDA,
            periodo=_PERIODO_RENDA,
            classificacoes=_CLASSIFICACOES_RENDA,
        )
        estimados = consulta.valores_estimados(MUNICIPIOS_ESTIMADOS)
        if estimados > self.configuracao.sidra_max_valores:
            raise SidraErro(tabela, f"{estimados} valores numa requisição excedem o limite")
        url = consulta.url(self.configuracao.sidra_base_url)
        conteudo, dados = self._consultar(url)
        _periodo_unico(dados, "D3C", url, esperado=_PERIODO_RENDA)
        linhas, parquet = self._gravar_sidra_unica(tabela, conteudo, list(dados[0]), url)
        entrada = EntradaIBGE(
            tabela, url, _PERIODO_RENDA, None, len(conteudo), None, 1, linhas, parquet,
            self.ingerido_em.isoformat(),
        )  # fmt: skip
        self._registrar(entrada, inicio)

    # -- REGIC

    def _arquivo_regic(self, nome: str, obtidos: dict[str, Path]) -> Path:
        """Baixa (ou localiza na origem local) um arquivo da REGIC, uma vez por execução."""
        if nome in obtidos:
            return obtidos[nome]
        if self.origem_local is not None:
            caminho = self._local(nome)
            if not caminho.is_file():
                raise PlanilhaIBGEErro(nome, f"ausente em {caminho.parent}")
        else:
            caminho = baixar_com_retentativas(
                self.http,
                self.configuracao.regic_base_url + nome,
                self.baixados / "regic" / nome,
                tamanho_esperado=None,
                tentativas=self.configuracao.tentativas,
                hosts_permitidos=self.configuracao.hosts_permitidos,
                dormir=self.dormir,
                max_retomadas=self.configuracao.max_retomadas,
                tempo_limite_total_s=self.configuracao.tempo_limite_total_s,
            )
        obtidos[nome] = caminho
        return caminho

    def regic(self, planilhas: Sequence[PlanilhaRegic] = PLANILHAS_REGIC) -> None:
        obtidos = self._regic_obtidos  # o zip de ligações e atração serve a duas tabelas
        for planilha in planilhas:
            tabela = planilha.tabela
            inicio = time.monotonic()
            anterior = self.anteriores.get(tabela)
            if self.origem_local is None:
                if not self.renovar and _reaproveitavel(self.configuracao, anterior):
                    self._pular(tabela, "planilha estática já no raw")
                    continue
                arquivo = self._arquivo_regic(planilha.arquivo, obtidos)
                sha = sha256_arquivo(arquivo)
                origem = self.configuracao.regic_base_url + planilha.arquivo
            else:
                if self._ausente_local(tabela, self._local(planilha.arquivo)):
                    continue
                arquivo = self._arquivo_regic(planilha.arquivo, obtidos)
                sha = sha256_arquivo(arquivo)
                origem = str(arquivo)
                if not self.renovar and _reaproveitavel(self.configuracao, anterior, sha256=sha):
                    self._pular(tabela, "arquivo local sem mudanças")
                    continue

            def gravar(
                con: duckdb.DuckDBPyConnection,
                destino: Path,
                trabalho: Path,
                planilha: PlanilhaRegic = planilha,
                arquivo: Path = arquivo,
            ) -> int:
                xlsx = arquivo
                if planilha.membro_zip is not None:
                    extraidos = extrair_zip_seguro(arquivo, trabalho / "zip")
                    candidatos = [p for p in extraidos if p.name == planilha.membro_zip]
                    if not candidatos:
                        raise PlanilhaIBGEErro(arquivo.name, f"sem {planilha.membro_zip}")
                    xlsx = candidatos[0]
                return _gravar_aba_parquet(
                    con, xlsx, planilha.aba, destino, trabalho, self.ingerido_em
                )

            linhas, parquet = self._publicar_parquet(tabela, gravar)
            entrada = EntradaIBGE(
                tabela, origem, _PERIODO_REGIC, arquivo.name, arquivo.stat().st_size, sha, 0,
                linhas, parquet, self.ingerido_em.isoformat(),
            )  # fmt: skip
            self._registrar(entrada, inicio)


def _periodo_unico(dados: list[dict], chave: str, onde: str, *, esperado: str | None = None) -> str:
    periodos = sorted({str(d.get(chave)) for d in dados})
    if len(periodos) != 1 or (esperado is not None and periodos[0] != esperado):
        alvo = f" (esperado {esperado})" if esperado else ""
        raise SidraErro(onde, f"períodos {periodos} na resposta{alvo}")
    return periodos[0]


# incremento: enriquecimento_ibge
def ingerir_ibge(
    configuracao: Configuracao,
    *,
    ingerido_em: datetime,
    origem_local: Path | None = None,
    renovar: bool = False,
    http: httpx.Client | None = None,
    dormir: Callable[[float], None] = time.sleep,
) -> tuple[EntradaIBGE, ...]:
    """Baixa (ou lê de `origem_local`) e converte as tabelas IBGE; devolve o manifesto IBGE.

    Tabelas já presentes no raw com o mesmo período/arquivo são puladas, salvo `renovar`. O
    manifesto `_manifestos/ibge/manifesto.json` é regravado a cada tabela concluída. Uma falha numa
    tabela não levanta (RIBG-02): avisa, deixa o raw anterior intacto (ou publica o vazio), registra
    a falha no manifesto e, no CEMPRE, deixa as fatias prontas no disco para a próxima execução. No
    fim, toda tabela ausente do raw é publicada vazia (`garantir_raw_ibge`).
    """
    cliente_http = http if http is not None else httpx.Client(timeout=configuracao.tempo_limite_s)
    try:
        ingestao = _Ingestao(
            configuracao,
            origem_local=origem_local,
            renovar=renovar,
            ingerido_em=ingerido_em,
            http=cliente_http,
            dormir=dormir,
        )
        try:
            periodo, categorias = ingestao.categorias_cempre()
        except _FALHAS_TOLERADAS as exc:
            ingestao.falhou("cempre_categorias_cnae", exc)
            if origem_local is None:  # na rede, o CEMPRE municipal depende do período e das classes
                ingestao.falhou("cempre_municipio_cnae", exc)
            else:
                _tolerar(
                    ingestao, "cempre_municipio_cnae", lambda: ingestao.cempre_municipio(None, [])
                )
        else:
            _tolerar(
                ingestao,
                "cempre_municipio_cnae",
                lambda: ingestao.cempre_municipio(periodo, categorias),
            )
        _tolerar(ingestao, "renda_censo_2022_municipio", ingestao.renda)
        for planilha in PLANILHAS_REGIC:
            _tolerar(ingestao, planilha.tabela, lambda p=planilha: ingestao.regic((p,)))
    finally:
        if http is None:
            cliente_http.close()
    garantir_raw_ibge(configuracao)
    return tuple(ingestao.entradas[nome] for nome in sorted(ingestao.entradas))


# RIBG-02: falhas do IBGE (rede, API, planilha, conversão) viram aviso; o mês segue
_FALHAS_TOLERADAS = (ErroIngestao, httpx.HTTPError, duckdb.Error, OSError)


def _tolerar(ingestao: _Ingestao, tabela: str, etapa: Callable[[], object]) -> None:
    try:
        etapa()
    except _FALHAS_TOLERADAS as exc:
        ingestao.falhou(tabela, exc)
